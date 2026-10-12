"""Generate the purpose-made evaluation documents in this folder.

All content is invented (a fictional company, Halvorsen Robotics), so the answer model cannot
know any answer in advance. Each document targets one kind of content the pipeline must read:
plain text, tables, charts, scanned pages, slide layouts, diagrams, and photographed text.

Run from the project root:  python eval_docs/make_eval_docs.py
"""
import io
import os

import docx
import pymupdf
from docx.shared import Inches
from PIL import Image, ImageDraw, ImageFilter, ImageFont

OUT = os.path.dirname(os.path.abspath(__file__))


def font(size: int):
    return ImageFont.load_default(size=size)


def png(img: Image.Image) -> bytes:
    buf = io.BytesIO()
    img.save(buf, format="PNG")
    return buf.getvalue()


def bar_chart(title: str, labels: list[str], values: list[float], unit: str) -> bytes:
    img = Image.new("RGB", (720, 420), "white")
    g = ImageDraw.Draw(img)
    g.text((40, 15), title, fill="black", font=font(24))
    top = max(values)
    for i, (label, value) in enumerate(zip(labels, values)):
        x = 70 + i * (600 // len(labels))
        h = int(300 * value / top)
        g.rectangle([x, 380 - h, x + 70, 380], fill=(70, 110, 170))
        g.text((x, 385), label, fill="black", font=font(16))
        g.text((x, 355 - h), f"{value:g}{unit}", fill="black", font=font(16))
    return png(img)


def floor_plan() -> bytes:
    img = Image.new("RGB", (900, 560), "white")
    g = ImageDraw.Draw(img)
    g.text((250, 12), "Oslo office - 3rd floor plan", fill="black", font=font(24))
    rooms = {
        "Fjord meeting room": (40, 60, 330, 290), "Server room": (330, 60, 590, 290),
        "Robot test lab": (590, 60, 860, 290), "Kitchen": (40, 290, 330, 520),
        "Open office": (330, 290, 590, 520), "Quiet room": (590, 290, 860, 520),
    }
    for name, box in rooms.items():
        g.rectangle(box, outline="black", width=4)
        g.text((box[0] + 14, box[1] + 14), name, fill="black", font=font(20))
    # First aid kit: a red cross inside the kitchen, labelled only by its symbol and a short tag.
    g.rectangle((250, 440, 300, 455), fill="red")
    g.rectangle((268, 422, 283, 473), fill="red")
    g.text((200, 480), "First aid", fill="red", font=font(16))
    return png(img)


def timeline() -> bytes:
    img = Image.new("RGB", (1000, 300), "white")
    g = ImageDraw.Draw(img)
    g.text((300, 15), "Picker X2 launch timeline", fill="black", font=font(26))
    stages = [("Design", "Feb"), ("Prototype", "Apr"), ("Pilot - Austin", "Jul"), ("General release", "Oct")]
    for i, (name, month) in enumerate(stages):
        x = 30 + i * 245
        g.rounded_rectangle((x, 100, x + 190, 190), radius=14, outline=(40, 90, 60), width=4)
        g.text((x + 14, 125), name, fill="black", font=font(20))
        g.text((x + 14, 205), month, fill=(40, 90, 60), font=font(22))
        if i < len(stages) - 1:
            g.line((x + 192, 145, x + 243, 145), fill="black", width=4)
            g.polygon([(x + 243, 145), (x + 231, 137), (x + 231, 153)], fill="black")
    return png(img)


def org_chart() -> bytes:
    img = Image.new("RGB", (1000, 460), "white")
    g = ImageDraw.Draw(img)
    g.text((380, 10), "Launch team", fill="black", font=font(26))
    boxes = {
        "ceo": ((370, 60, 630, 120), "Ingrid Halvorsen", "CEO"),
        "vpp": ((120, 190, 400, 250), "Marcus Lee", "VP Product"),
        "vpe": ((600, 190, 880, 250), "Sofia Brandt", "VP Engineering"),
        "lm": ((120, 330, 400, 390), "Aiko Tanaka", "Launch Manager"),
        "rl": ((600, 330, 880, 390), "Ravi Patel", "Robotics Lead"),
    }
    for (x0, y0, x1, y1), name, role in boxes.values():
        g.rectangle((x0, y0, x1, y1), outline="black", width=3)
        g.text((x0 + 12, y0 + 6), name, fill="black", font=font(20))
        g.text((x0 + 12, y0 + 32), role, fill=(90, 90, 90), font=font(16))
    for parent, child in [("ceo", "vpp"), ("ceo", "vpe"), ("vpp", "lm"), ("vpe", "rl")]:
        (px0, _, px1, py1), _, _ = boxes[parent]
        (cx0, cy0, cx1, _), _, _ = boxes[child]
        g.line(((px0 + px1) // 2, py1, (cx0 + cx1) // 2, cy0), fill="black", width=3)
    return png(img)


def annual_report() -> None:
    doc = pymupdf.open()

    def page_with(blocks):
        page = doc.new_page()
        y = 60
        for kind, content in blocks:
            if kind == "h1":
                page.insert_text((60, y + 24), content, fontsize=24, fontname="hebo")
                y += 50
            elif kind == "h2":
                page.insert_text((60, y + 18), content, fontsize=17, fontname="hebo")
                y += 36
            elif kind == "p":
                rect = pymupdf.Rect(60, y, 540, y + 120)
                page.insert_textbox(rect, content, fontsize=11, fontname="helv")
                y += 14 * (len(content) // 80 + 2)
            elif kind == "img":
                page.insert_image(pymupdf.Rect(80, y, 520, y + 256), stream=content)
                y += 270

    page_with([
        ("h1", "Halvorsen Robotics - Annual Report 2025"),
        ("h2", "Overview"),
        ("p", "Halvorsen Robotics builds autonomous picking robots for warehouses. The company was "
              "founded in 2014 by Ingrid Halvorsen, who is still its CEO. At the end of 2025 it "
              "employed 412 people across two offices, in Oslo and Austin."),
        ("h2", "Financial results"),
        ("p", "Total revenue for 2025 was 48.6 million dollars, up 18 percent from 2024. Revenue by "
              "quarter is shown in the chart below."),
        ("img", bar_chart("Revenue by quarter, 2025 ($M)", ["Q1", "Q2", "Q3", "Q4"], [10.2, 11.5, 13.9, 13.0], "")),
    ])
    page_with([
        ("h2", "People"),
        ("p", "The chart below shows headcount by department at the end of the year."),
        ("img", bar_chart("Headcount by department", ["Engineering", "Operations", "Sales", "Admin"], [198, 104, 62, 48], "")),
        ("h2", "Outlook"),
        ("p", "In 2026 the company plans to open a third office, in Singapore, and targets revenue of "
              "60 million dollars. The 2026 IT hardware budget is 1.2 million dollars."),
    ])
    doc.save(os.path.join(OUT, "annual_report.pdf"), garbage=4, deflate=True)


def employee_handbook() -> None:
    d = docx.Document()
    d.add_heading("Halvorsen Robotics Employee Handbook", 0)
    d.add_heading("Time off", 1)
    d.add_paragraph("Vacation days depend on how long you have worked at the company.")
    t = d.add_table(rows=1, cols=2)
    t.rows[0].cells[0].text, t.rows[0].cells[1].text = "Years of service", "Vacation days per year"
    for years, days in [("0-4 years", "20"), ("5-9 years", "25"), ("10 or more years", "30")]:
        row = t.add_row().cells
        row[0].text, row[1].text = years, days
    d.add_heading("Travel", 1)
    d.add_paragraph("Daily allowances and hotel limits for business travel:")
    t = d.add_table(rows=1, cols=3)
    for cell, text in zip(t.rows[0].cells, ["Role", "Daily allowance", "Hotel limit per night"]):
        cell.text = text
    for role, allowance, hotel in [("Engineer", "$60", "$180"), ("Manager", "$85", "$250"), ("Director", "$110", "$320")]:
        row = t.add_row().cells
        row[0].text, row[1].text, row[2].text = role, allowance, hotel
    d.add_heading("Oslo office layout", 1)
    d.add_paragraph("The plan below shows the 3rd floor of the Oslo office.")
    d.add_picture(io.BytesIO(floor_plan()), width=Inches(6))
    d.add_heading("Equipment", 1)
    d.add_paragraph("Company laptops are replaced every 3 years.")
    d.save(os.path.join(OUT, "employee_handbook.docx"))


def scanned_invoice() -> None:
    # Rendered as a picture and placed on a PDF page with no text layer, like a scan.
    img = Image.new("RGB", (1240, 1600), (250, 249, 245))
    g = ImageDraw.Draw(img)
    lines = [
        ("INVOICE  #INV-2025-0871", 40), ("", 20), ("Date: 12 November 2025", 28),
        ("From: Nordic Servo Supply AS, Bergen", 28), ("Bill to: Halvorsen Robotics, Oslo", 28), ("", 20),
        ("Item                              Qty    Unit price    Amount", 26),
        ("Servo motor SM-220                 40      $310        $12,400", 26),
        ("Controller board CB-9              15      $540         $8,100", 26),
        ("Shipping                                                 $650", 26), ("", 20),
        ("TOTAL DUE:  $21,150", 36), ("Payment due within 30 days (by 12 December 2025)", 26),
    ]
    y = 90
    for text, size in lines:
        g.text((90, y), text, fill=(30, 30, 30), font=font(size))
        y += size + 30
    img = img.rotate(-0.8, expand=False, fillcolor=(250, 249, 245)).filter(ImageFilter.GaussianBlur(0.6))
    doc = pymupdf.open()
    page = doc.new_page()
    page.insert_image(page.rect, stream=png(img))
    doc.save(os.path.join(OUT, "scanned_invoice.pdf"), garbage=4, deflate=True)


def launch_deck() -> None:
    doc = pymupdf.open()
    slides = [
        ("Picker X2 Launch Plan", ["Product team, January 2026"], None),
        ("Key features", ["Picks up to 1,200 items per hour", "Battery life: 14 hours",
                          "Price: 89,000 dollars per unit"], None),
        ("Launch timeline", [], timeline()),
        ("Team", [], org_chart()),
    ]
    for title, bullets, image in slides:
        page = doc.new_page(width=842, height=595)
        page.insert_text((50, 80), title, fontsize=34, fontname="hebo")
        for i, bullet in enumerate(bullets):
            page.insert_text((70, 160 + i * 50), f"- {bullet}", fontsize=20, fontname="helv")
        if image:
            page.insert_image(pymupdf.Rect(50, 130, 790, 560), stream=image)
    doc.save(os.path.join(OUT, "launch_deck.pdf"), garbage=4, deflate=True)


def whiteboard() -> None:
    img = Image.new("RGB", (1100, 700), (236, 238, 240))
    g = ImageDraw.Draw(img)
    lines = [
        ("SPRINT 14 RETRO", 44), ("", 10),
        ("Went well: new gripper calibration -> 30% fewer dropped items", 28),
        ("Problem: robot test lab is booked out every Thursday", 28),
        ("Action: Ravi books the test lab for Mondays", 28),
        ("Next demo: 6 February", 28),
    ]
    y = 70
    for text, size in lines:
        g.text((60, y), text, fill=(25, 60, 160), font=font(size))
        y += size + 45
    img = img.rotate(2.0, expand=True, fillcolor=(120, 120, 120)).filter(ImageFilter.GaussianBlur(0.8))
    img.save(os.path.join(OUT, "sprint_whiteboard.jpg"), quality=80)


def it_faq() -> None:
    with open(os.path.join(OUT, "it_faq.md"), "w", encoding="utf-8") as f:
        f.write(
            "# IT FAQ\n\n"
            "## Passwords\n"
            "Passwords must be at least 14 characters long and must be changed every 180 days.\n\n"
            "## VPN\n"
            "Connect through the HalvorsenSecure VPN. For VPN problems, call IT support on extension 4410.\n\n"
            "## Laptops\n"
            "Request a replacement laptop through the IT portal. Laptops are replaced every 3 years.\n"
        )


if __name__ == "__main__":
    annual_report()
    employee_handbook()
    scanned_invoice()
    launch_deck()
    whiteboard()
    it_faq()
    print("Wrote:", ", ".join(sorted(f for f in os.listdir(OUT) if not f.endswith(".py"))))
